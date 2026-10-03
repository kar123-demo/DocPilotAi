from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _heading(document: Document, value: str, level: int = 1) -> None:
    document.add_heading(value, level=level)


def _code_block(document: Document, value: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(value)
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def _bullets(document: Document, items: list[str]) -> None:
    for item in items or []:
        document.add_paragraph(str(item), style="List Bullet")


def _key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value


def build_report(output_path: Path, project_name: str, repository_url: str,
                  narrative: dict[str, str], structural_uml: str, diagram_path: Path | None,
                  behavioral_uml: str | None = None, behavior_diagram_path: Path | None = None,
                  title_page: dict[str, Any] | None = None, sections: dict[str, Any] | None = None) -> None:
    title_page = title_page or {}
    sections = sections or {}
    document = Document()
    section = document.sections[0]
    section.top_margin, section.bottom_margin = Inches(0.75), Inches(0.75)

    # ---- Title page --------------------------------------------------------------------
    title = document.add_heading(project_name, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Project Report").alignment = WD_ALIGN_PARAGRAPH.CENTER
    for label, value in (
        ("College/Organization", title_page.get("college_name", "Not specified")),
        ("Guide/Mentor", title_page.get("guide_name", "Not specified")),
        ("Repository", repository_url),
    ):
        paragraph = document.add_paragraph(f"{label}: {value}")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    team_members = title_page.get("team_members") or []
    if team_members:
        paragraph = document.add_paragraph("Team Members: " + ", ".join(team_members))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    toc = [
        "Abstract",
        "Introduction",
        "Literature Survey / Existing System",
        "Structural and Behavioral Diagrams",
        "Technology Stack",
        "Limitations",
        "Conclusion",
    ]
    _heading(document, "Table of Contents")
    for number, entry in enumerate(toc, start=1):
        document.add_paragraph(f"{number}. {entry}")
    document.add_page_break()

    # ---- Abstract -----------------------------------------------------------------------
    _heading(document, "Abstract")
    document.add_paragraph(narrative.get("abstract", ""))

    # ---- Introduction ---------------------------------------------------------------------
    _heading(document, "Introduction")
    document.add_paragraph(narrative.get("introduction", ""))

    # ---- Literature survey -------------------------------------------------------------
    _heading(document, "Literature Survey / Existing System")
    document.add_paragraph(narrative.get("literature_survey", ""))

    # ---- Structural and behavioral diagrams ----------------------------------------------
    _heading(document, "Structural and Behavioral Diagrams")
    diagrams = sections.get("diagrams")
    if diagrams:
        for category, entries in (("Structural UML", diagrams.get("structural", [])),
                                  ("Behavioral UML", diagrams.get("behavioral", []))):
            _heading(document, category, level=2)
            for diagram in entries:
                _heading(document, diagram["title"], level=3)
                image = Path(diagram.get("image_path", ""))
                if image.is_file():
                    document.add_picture(str(image), width=Inches(6.2))
                else:
                    _code_block(document, diagram["code"])
    else:
        _heading(document, "Structural Diagram", level=2)
        if diagram_path and diagram_path.exists():
            document.add_picture(str(diagram_path), width=Inches(6.2))
        else:
            document.add_paragraph("Diagram image rendering was unavailable. PlantUML source follows.")
            _code_block(document, structural_uml)
        if behavior_diagram_path and behavior_diagram_path.exists():
            _heading(document, "Behavioral Diagram", level=2)
            document.add_picture(str(behavior_diagram_path), width=Inches(6.2))
        elif behavioral_uml:
            _heading(document, "Behavioral Diagram", level=2)
            _code_block(document, behavioral_uml)

    # ---- Technology stack -----------------------------------------------------------------
    _heading(document, "Technology Stack")
    stack = sections.get("technology_stack", {})
    for label, key in (("Frontend", "frontend"), ("Backend", "backend"), ("Database", "database"), ("Tools/Libraries", "tools")):
        _heading(document, label, level=2)
        _bullets(document, stack.get(key, []))

    # ---- Limitations -----------------------------------------------------------------------
    _heading(document, "Limitations")
    _bullets(document, sections.get("limitations", []))

    # ---- Conclusion -------------------------------------------------------------------------
    _heading(document, "Conclusion")
    document.add_paragraph(narrative.get("conclusion", ""))

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Aptos"
            if run.font.size is None:
                run.font.size = Pt(11)
    document.save(output_path)
